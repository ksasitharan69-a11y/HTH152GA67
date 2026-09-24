import os
import re
import logging
from typing import Optional
from pypdf import PdfReader
import docx

logger = logging.getLogger(__name__)

class ResumeParser:
    """Extracts, cleans, and normalizes text from PDF, DOCX, and text files."""

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """Extract text content from a PDF file using pypdf."""
        extracted_pages = []
        try:
            reader = PdfReader(file_path)
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    extracted_pages.append(page_text.strip())
            return "\n\n".join(extracted_pages)
        except Exception as e:
            logger.warning(f"pypdf reader error for {file_path}: {e}")
            raise ValueError(f"Failed to extract text from PDF: {str(e)}")

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text content from a DOCX file using python-docx."""
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract text inside tables
            table_cells = []
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        table_cells.append(" | ".join(row_texts))
                        
            all_text = "\n".join(paragraphs)
            if table_cells:
                all_text += "\n\n" + "\n".join(table_cells)
            return all_text
        except Exception as e:
            logger.warning(f"docx reader error for {file_path}: {e}")
            raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

    @classmethod
    def extract_text(cls, file_path: str, mime_type: Optional[str] = None) -> str:
        """Extract and normalize text based on file extension and/or mime type with safe fallback."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found at path: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()
        raw_text = ""

        if ext == ".pdf" or (mime_type and "pdf" in mime_type.lower()):
            try:
                raw_text = cls.extract_text_from_pdf(file_path)
            except Exception as e:
                # Safe fallback: check if file content is readable text (e.g. text stream with .pdf extension)
                try:
                    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                        fallback_text = f.read()
                    if len(fallback_text.strip()) > 10:
                        raw_text = fallback_text
                    else:
                        raise e
                except Exception:
                    raise e
        elif ext in [".docx", ".doc"] or (mime_type and "word" in mime_type.lower()):
            try:
                raw_text = cls.extract_text_from_docx(file_path)
            except Exception as e:
                try:
                    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                        fallback_text = f.read()
                    if len(fallback_text.strip()) > 10:
                        raw_text = fallback_text
                    else:
                        raise e
                except Exception:
                    raise e
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                raw_text = f.read()
        else:
            raise ValueError(f"Unsupported file format '{ext}'. Supported formats: PDF, DOCX, TXT")

        return cls.normalize_text(raw_text)

    @staticmethod
    def normalize_text(text: str) -> str:
        """Clean and normalize extracted text while preserving section structure."""
        if not text:
            return ""

        # Normalize line endings
        text = text.replace("\r\n", "\n").replace("\r", "\n")

        # Replace excessive horizontal whitespace with a single space
        text = re.sub(r"[ \t]+", " ", text)

        # Consolidate 3+ newlines to 2 newlines (preserve paragraphs)
        text = re.sub(r"\n{3,}", "\n\n", text)

        # Strip unprintable control characters except standard tabs/newlines
        text = "".join(ch for ch in text if ch.isprintable() or ch in "\n\t")

        return text.strip()
